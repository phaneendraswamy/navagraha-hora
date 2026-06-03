from html import escape
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from backend.config import get_settings


SECTION_ORDER = ["Profile Summary", "Core Skills", "Professional Experience", "Projects", "Education", "Strengths", "Certifications"]


def _ensure_exports() -> Path:
    path = Path(get_settings().exports_dir)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _contact_line(personal: dict) -> str:
    return " | ".join(filter(None, [
        personal.get("email"),
        personal.get("phone"),
        personal.get("location"),
        personal.get("linkedin_url"),
        personal.get("github_url"),
        personal.get("portfolio_url"),
    ]))


def _skill_lines(data: dict) -> list[str]:
    groups = data.get("skill_groups") or {}
    if not groups:
        return [", ".join(data.get("skills", []))]
    return [f"**{category}:** {', '.join(values)}" for category, values in groups.items() if values]


def _clean_institution(value: str | None) -> str:
    if not value or value.lower().strip() == "extracted from profile":
        return ""
    return value


def build_preview(data: dict) -> str:
    personal = data["personal"]
    lines = [
        f"# {personal['full_name']}",
        _contact_line(personal),
        "",
        "## Profile Summary",
        data["summary"],
        "",
        "## Core Skills",
        *_skill_lines(data),
        "",
        "## Professional Experience",
    ]
    for exp in data["experiences"]:
        title = f"**{exp['role']} - {exp['company']}**"
        if exp.get("duration"):
            title += f" | {exp['duration']}"
        lines.append(title)
        lines.extend(f"- {bullet}" for bullet in exp["bullets"][:4])
        lines.append("")
    lines.append("## Projects")
    for project in data["projects"]:
        tech = ", ".join(project.get("technologies", [])[:10])
        lines.append(f"**{project['name']}**" + (f" | {tech}" if tech else ""))
        lines.extend(f"- {bullet}" for bullet in project["bullets"][:3])
        lines.append("")
    if data.get("education"):
        lines.append("## Education")
        for edu in data["education"]:
            detail = f"- {edu['degree']}"
            institution = _clean_institution(edu.get("institution"))
            if institution:
                detail += f" - {institution}"
            if edu.get("duration"):
                detail += f" | {edu['duration']}"
            lines.append(detail)
        lines.append("")
    if data.get("strengths"):
        lines.append("## Strengths")
        lines.extend(f"- {strength}" for strength in data["strengths"])
        lines.append("")
    if data.get("certifications"):
        lines.append("## Certifications")
        lines.extend(f"- {cert['name']}" + (f" - {cert.get('organization')}" if cert.get("organization") else "") for cert in data["certifications"])
    return "\n".join(lines)


def export_docx(data: dict, template_name: str) -> str:
    path = _ensure_exports() / f"resume_{uuid4().hex[:10]}.docx"
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.2)

    personal = data["personal"]
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_p.add_run(personal["full_name"].upper())
    name_run.bold = True
    name_run.font.size = Pt(16)
    name_run.font.color.rgb = RGBColor(15, 23, 42)
    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_run = contact_p.add_run(_contact_line(personal))
    contact_run.font.size = Pt(8.8)

    def heading(text: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(10.2)
        r.font.color.rgb = RGBColor(15, 23, 42)
        p.paragraph_format.left_indent = Inches(0)

    def bullet(text: str) -> None:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        p.add_run(text)

    heading("Profile Summary")
    doc.add_paragraph(data["summary"])

    heading("Core Skills")
    groups = data.get("skill_groups") or {}
    if groups:
        for category, values in groups.items():
            p = doc.add_paragraph()
            p.add_run(f"{category}: ").bold = True
            p.add_run(", ".join(values))
    else:
        doc.add_paragraph(", ".join(data.get("skills", [])))

    heading("Professional Experience")
    for exp in data["experiences"]:
        p = doc.add_paragraph()
        p.add_run(f"{exp['role']} - {exp['company']}").bold = True
        if exp.get("duration"):
            p.add_run(f" | {exp['duration']}")
        for item in exp["bullets"][:4]:
            bullet(item)

    heading("Projects")
    for project in data["projects"]:
        p = doc.add_paragraph()
        p.add_run(project["name"]).bold = True
        tech = ", ".join(project.get("technologies", [])[:10])
        if tech:
            p.add_run(f" | {tech}")
        for item in project["bullets"][:3]:
            bullet(item)

    if data.get("education"):
        heading("Education")
        for edu in data["education"]:
            line = edu["degree"]
            institution = _clean_institution(edu.get("institution"))
            if institution:
                line += f" - {institution}"
            if edu.get("duration"):
                line += f" | {edu['duration']}"
            bullet(line)

    if data.get("strengths"):
        heading("Strengths")
        for strength in data["strengths"]:
            bullet(strength)

    if data.get("certifications"):
        heading("Certifications")
        for cert in data["certifications"]:
            bullet(f"{cert['name']}" + (f" - {cert.get('organization')}" if cert.get("organization") else ""))

    doc.save(path)
    return str(path)


def _plain_skill_lines(data: dict) -> list[str]:
    groups = data.get("skill_groups") or {}
    if groups:
        return [f"<b>{escape(category)}:</b> {escape(', '.join(values))}" for category, values in groups.items()]
    return [escape(", ".join(data.get("skills", [])))]


def export_pdf(data: dict, template_name: str) -> str:
    path = _ensure_exports() / f"resume_{uuid4().hex[:10]}.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=32, leftMargin=32, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    body = ParagraphStyle("BodyResume", parent=styles["BodyText"], fontSize=8.8, leading=11.2, spaceAfter=2)
    heading = ParagraphStyle("HeadingResume", parent=styles["Heading4"], fontSize=9.8, leading=12, textColor=colors.HexColor("#0f172a"), spaceBefore=7, spaceAfter=3)
    title = ParagraphStyle("TitleResume", parent=styles["Title"], fontSize=17, leading=20, alignment=1, spaceAfter=2)
    contact_style = ParagraphStyle("ContactResume", parent=body, alignment=1, fontSize=8.2, leading=10)
    story = []

    personal = data["personal"]
    story.extend([
        Paragraph(f"<b>{escape(personal['full_name'].upper())}</b>", title),
        Paragraph(escape(_contact_line(personal)), contact_style),
        Spacer(1, 6),
    ])

    def add_section(title_text: str, content: list[str]) -> None:
        if not content:
            return
        story.append(Paragraph(f"<b>{escape(title_text.upper())}</b>", heading))
        for item in content:
            story.append(Paragraph(item, body))
        story.append(Spacer(1, 3))

    add_section("Profile Summary", [escape(data["summary"])])
    add_section("Core Skills", _plain_skill_lines(data))

    exp_lines = []
    for exp in data["experiences"]:
        exp_title = f"<b>{escape(exp['role'])} - {escape(exp['company'])}</b>"
        if exp.get("duration"):
            exp_title += f" | {escape(exp['duration'])}"
        exp_lines.append(exp_title)
        exp_lines.extend(f"- {escape(item)}" for item in exp["bullets"][:4])
    add_section("Professional Experience", exp_lines)

    project_lines = []
    for project in data["projects"]:
        tech = ", ".join(project.get("technologies", [])[:10])
        title_line = f"<b>{escape(project['name'])}</b>"
        if tech:
            title_line += f" | {escape(tech)}"
        project_lines.append(title_line)
        project_lines.extend(f"- {escape(item)}" for item in project["bullets"][:3])
    add_section("Projects", project_lines)

    if data.get("education"):
        education = []
        for edu in data["education"]:
            line = edu["degree"]
            institution = _clean_institution(edu.get("institution"))
            if institution:
                line += f" - {institution}"
            if edu.get("duration"):
                line += f" | {edu['duration']}"
            education.append(escape(line))
        add_section("Education", education)
    if data.get("strengths"):
        add_section("Strengths", [f"- {escape(strength)}" for strength in data["strengths"]])
    if data.get("certifications"):
        add_section("Certifications", [
            escape(f"{cert['name']}" + (f" - {cert.get('organization')}" if cert.get("organization") else ""))
            for cert in data["certifications"]
        ])

    doc.build(story)
    return str(path)
